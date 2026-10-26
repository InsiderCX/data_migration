"""
Build the Ops validation workbook from the two CSV exports of sql/08_ops_lists.sql.

Usage:
  python3 -I scripts/3_build_ops_workbook.py invites.csv responses.csv out.xlsx [--bi Joined.csv]
          [--title "Stomatologija i estetika"] [--legacy-from 2023-05 --legacy-to 2024-10] [--source "EU staging, 26 Oct 2026"]

Sheets: Read me, Summary (COUNTIFS formulas over the data sheets), Invites sent, Responses.
Estimated cells yellow, legacy months grey. No names, phones or comments (the SQL does not export them).
--bi: the current BI system's Alchemer export (column "Response ID"); fills "In current BI system".
Afterwards recalculate so the formulas carry values:  python3 <xlsx skill>/scripts/recalc.py out.xlsx 120
"""
import argparse
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument("invites"); ap.add_argument("responses"); ap.add_argument("out")
ap.add_argument("--bi"); ap.add_argument("--title", default="Stomatologija i estetika")
ap.add_argument("--legacy-from", default="2023-05"); ap.add_argument("--legacy-to", default="2024-10")
ap.add_argument("--source", default="")
a = ap.parse_args()

inv = pd.read_csv(a.invites, dtype=str, keep_default_na=False)
resp = pd.read_csv(a.responses, dtype=str, keep_default_na=False)
if a.bi:
    bi_ids = set(pd.read_csv(a.bi, dtype=str, encoding="utf-8-sig", keep_default_na=False)["Response ID"])
    resp["in_bi"] = resp.alchemer_response_id.map(lambda x: "yes" if x in bi_ids else "no")
for df in (inv, resp):
    df["score"] = pd.to_numeric(df["score"], errors="coerce").astype("Int64")

INV_COLS = [("invite_id", "Invite ID (database)"), ("legacy_invite_id", "Legacy invite ID (sguid)"), ("source", "Source"),
            ("sent_local", "Sent (local time)"), ("sent_date_is", "Sent date is"), ("status", "Delivery status"),
            ("status_is", "Status is"), ("visit_local", "Visit date"), ("visit_date_is", "Visit date is"), ("doctor", "Doctor"),
            ("procedure", "Procedure"), ("procedure_is", "Procedure is"), ("responded", "Responded"),
            ("alchemer_response_id", "Alchemer response ID"), ("score", "Score"), ("response_status", "Response status"),
            ("survey_started_local", "Survey started (local time)"), ("send_month", "Send month")]
RESP_COLS = [("alchemer_response_id", "Alchemer response ID"), ("sguid", "Survey link ID (sguid)"), ("kind", "Type"),
             ("status", "Response status"), ("score", "Score"), ("started_local", "Survey started (local time)"),
             ("submitted_local", "Submitted (local time)"), ("language", "Language"), ("shown", "Shown in dashboards"),
             ("reason", "Why not shown"), ("invite_id_after_import", "Invite ID (database)"), ("in_bi", "In current BI system"),
             ("start_month", "Start month"), ("comment_sentiment", "Comment sentiment")]

F = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True)
FH = Font(name="Arial", size=10, bold=True, color="FFFFFF"); FI = Font(name="Arial", size=9, italic=True)
HDR = PatternFill("solid", fgColor="1F3A5F"); EST = PatternFill("solid", fgColor="FFF2CC"); GREY = PatternFill("solid", fgColor="F2F2F2")
thin = Border(bottom=Side(style="thin", color="BFBFBF"))


def cell(v):
    if v is pd.NA or v == "" or (isinstance(v, float) and pd.isna(v)):
        return None
    return int(v) if not isinstance(v, str) else v


def data_sheet(ws, df, cols, widths, est_cols=()):
    ws.append([h for _, h in cols])
    for c in ws[1]:
        c.font = FH; c.fill = HDR; c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in df[[k for k, _ in cols]].itertuples(index=False):
        ws.append([cell(v) for v in row])
    keys = [k for k, _ in cols]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F
        for k in est_cols:
            c = row[keys.index(k)]
            if c.value in ("estimated", "derived from comment"):
                c.fill = EST
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
rm = wb.active; rm.title = "Read me"
n_leg = (inv.source == "legacy").sum(); n_pl = (inv.source == "platform").sum()
for text, font in [
    (f"{a.title}: all survey invites sent and all survey responses", FB),
    (f"Source: {a.source}" if a.source else "Source: database export (sql/08_ops_lists.sql)", F), ("", F),
    ("Sheets", FB),
    ("Summary: per month, invites sent, delivered, responses, response rate (by send month) and responses / NPS (by survey start). Formulas over the two data sheets.", F),
    (f"Invites sent: every initial survey invite, {n_leg:,} legacy (imported) + {n_pl:,} platform, with its response if there is one.", F),
    (f"Responses: every survey response in the database ({len(resp):,}), including the ones dashboards cannot show and why.", F), ("", F),
    ("Real vs estimated", FB),
    ("Real: every response (score, status, dates), doctor, which invite a legacy response belongs to, all platform data.", F),
    ("Estimated (legacy only, yellow): which visits got an invite, send date, visit date, delivery status of invites without a response.", F),
    ("Derived (legacy only, yellow): procedure, from the visit comment.", F), ("", F),
    ("Definitions (same as Clinicdash)", FB),
    ("Response rate = responses / delivered invites, by the month the invite was sent (Message performance).", F),
    ("NPS = (promoters 9-10 minus detractors 0-6) / scored responses x 100, by the month the survey was started (NPS Trends). Partial responses with a score count.", F),
    ("No patient names, phone numbers or comments are included. Match rows by Alchemer response ID or legacy invite ID (sguid).", F),
]:
    rm.append([text]); rm.cell(rm.max_row, 1).font = font
rm.column_dimensions["A"].width = 140

wi = wb.create_sheet("Invites sent")
data_sheet(wi, inv, INV_COLS, [38, 14, 10, 19, 12, 15, 24, 17, 12, 20, 14, 18, 10, 12, 8, 13, 19, 10],
           est_cols=("sent_date_is", "status_is", "visit_date_is", "procedure_is"))
wr = wb.create_sheet("Responses")
data_sheet(wr, resp, RESP_COLS, [12, 26, 12, 13, 8, 19, 19, 11, 11, 26, 38, 12, 10, 13])
for row in wr.iter_rows(min_row=2):
    if row[8].value == "no":
        for c in row:
            c.fill = GREY

ws = wb.create_sheet("Summary", 1)
ni, nr = len(inv) + 1, len(resp) + 1
I = lambda c: f"'Invites sent'!${c}$2:${c}${ni}"
R = lambda c: f"Responses!${c}$2:${c}${nr}"
ws.append(["Month", "Legacy invites (estimated)", "Platform invites", "Invites sent", "Delivered", "Responses to these invites",
           "Response rate", "Responses (by survey start)", "Promoters", "Detractors", "NPS"])
for c in ws[1]:
    c.font = FH; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 42
first = min(m for m in list(inv.send_month) + list(resp.start_month) if m)
last_m = max(m for m in list(inv.send_month) + list(resp.start_month) if m)
months = list(pd.period_range(first, last_m, freq="M").strftime("%Y-%m"))
for i, m in enumerate(months, start=2):
    ws.append([m, f'=COUNTIFS({I("C")},"legacy",{I("R")},$A{i})', f'=COUNTIFS({I("C")},"platform",{I("R")},$A{i})',
               f"=B{i}+C{i}", f'=COUNTIFS({I("R")},$A{i},{I("F")},"DELIVERED")', f'=COUNTIFS({I("R")},$A{i},{I("M")},"yes")',
               f"=IF(E{i}=0,0,F{i}/E{i})", f'=COUNTIFS({R("M")},$A{i},{R("I")},"yes")',
               f'=COUNTIFS({R("M")},$A{i},{R("I")},"yes",{R("E")},">=9")', f'=COUNTIFS({R("M")},$A{i},{R("I")},"yes",{R("E")},"<=6")',
               f'=IF(H{i}=0,"",ROUND((I{i}-J{i})/H{i}*100,1))'])
last = len(months) + 1; t = last + 1; lp = t + 1
ws.append(["Total"] + [f"=SUM({c}2:{c}{last})" for c in "BCDEF"] + [f"=IF(E{t}=0,0,F{t}/E{t})"]
          + [f"=SUM({c}2:{c}{last})" for c in "HIJ"] + [f'=IF(H{t}=0,"",ROUND((I{t}-J{t})/H{t}*100,1))'])
ws.append([f"Legacy period ({a.legacy_from} to {a.legacy_to})", f'=COUNTIFS({I("C")},"legacy")', 0, f"=B{lp}+C{lp}",
           f'=COUNTIFS({I("C")},"legacy",{I("F")},"DELIVERED")', f'=COUNTIFS({I("C")},"legacy",{I("M")},"yes")',
           f"=IF(E{lp}=0,0,F{lp}/E{lp})", f'=COUNTIFS({R("C")},"legacy",{R("I")},"yes")',
           f'=COUNTIFS({R("C")},"legacy",{R("I")},"yes",{R("E")},">=9")', f'=COUNTIFS({R("C")},"legacy",{R("I")},"yes",{R("E")},"<=6")',
           f'=IF(H{lp}=0,"",ROUND((I{lp}-J{lp})/H{lp}*100,1))'])
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.font = F; c.border = thin
    r[6].number_format = "0.0%"; r[10].number_format = "0.0"
for r in (t, lp):
    for c in ws[r]:
        c.font = FB
for i, m in enumerate(months, start=2):
    if a.legacy_from <= m <= a.legacy_to:
        for c in ws[i]:
            c.fill = GREY
        ws.cell(i, 2).fill = EST
ws.cell(lp, 2).fill = EST
ws.cell(lp + 2, 1, "Grey rows = legacy period. Yellow = estimated. Response rate counts responses by the month the invite was sent; "
        "NPS by the month the survey was started (as Clinicdash does), so the two response columns differ slightly.").font = FI
ws.column_dimensions["A"].width = 34
for col in "BCDEFGHIJK":
    ws.column_dimensions[col].width = 14
ws.freeze_panes = "B2"
wb.save(a.out)
print(f"wrote {a.out}: {len(inv)} invites, {len(resp)} responses, {len(months)} months. Now run recalc.py on it.")
